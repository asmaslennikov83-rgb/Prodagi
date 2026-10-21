from __future__ import annotations

import math
import re
from decimal import Decimal, ROUND_CEILING
from collections import Counter, defaultdict
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from report_builder import Product


@dataclass
class Bundle:
    name: str
    barcode: str
    components: list[str]


@dataclass
class PurchaseRow:
    product: Product
    bundle_sales_units: int = 0
    fbo_direct_stock: int = 0
    fbs_direct_stock: int = 0
    fbo_stock_inside_bundles: int = 0
    fbs_stock_inside_bundles: int = 0
    selected_stock: int = 0
    average_per_day: float = 0.0
    required_stock: int = 0
    purchase_qty: int = 0
    note: str = ''
    cabinet_fbo: dict[str, int] = field(default_factory=dict)
    cabinet_required: dict[str, int] = field(default_factory=dict)
    cabinet_needs: dict[str, int] = field(default_factory=dict)
    total_need: int = 0
    cabinet_days: dict[str, int | None] = field(default_factory=dict)
    cabinet_average: dict[str, float] = field(default_factory=dict)

    @property
    def total_consumption(self) -> int:
        return self.product.fbo + self.product.fbs + self.bundle_sales_units

    @property
    def fbo_physical_stock(self) -> int:
        return self.fbo_direct_stock + self.fbo_stock_inside_bundles

    @property
    def fbs_physical_stock(self) -> int:
        return self.fbs_direct_stock + self.fbs_stock_inside_bundles

    @property
    def physical_stock(self) -> int:
        return self.selected_stock

    @property
    def days_remaining(self) -> float:
        """How many days the selected stock will last at the current average order rate."""
        if self.average_per_day <= 0:
            return 0.0
        return self.physical_stock / self.average_per_day


@dataclass
class BundleAuditRow:
    bundle_name: str
    bundle_barcode: str
    fbo: int
    fbs: int
    bundle_orders: int
    bundle_stock_fbo: int
    bundle_stock_fbs: int
    bundle_stock_selected: int
    component_barcode: str
    component_name: str
    qty_in_bundle: int
    component_sales_consumption: int
    component_stock_inside_bundles_fbo: int
    component_stock_inside_bundles_fbs: int
    component_stock_inside_bundles_selected: int


@dataclass
class ProcurementResult:
    rows: list[PurchaseRow]
    bundle_rows: list[BundleAuditRow]
    errors: list[tuple[str, str, str]] = field(default_factory=list)
    source_products: list[Product] = field(default_factory=list)
    warnings: list[tuple[str, str, str]] = field(default_factory=list)


def _norm_header(value) -> str:
    text = '' if value is None else str(value)
    text = text.strip().casefold().replace('ё', 'е')
    return re.sub(r'[^a-zа-я0-9]+', '', text)


def barcode_text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, bool):
        return ''
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, '.15g').strip()
    text = str(value).strip()
    if not text:
        return ''
    if re.fullmatch(r'\d+\.0+', text):
        return text.split('.', 1)[0]
    return re.sub(r'\s+', '', text)


def int_qty(value) -> int:
    if value in (None, ''):
        return 0
    if isinstance(value, str):
        value = value.replace('\xa0', '').replace(' ', '').replace(',', '.')
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def _read_xlsx_rows(path: str | Path) -> list[list]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _read_xls_rows(path: str | Path) -> list[list]:
    try:
        import xlrd  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            'Для чтения .xls не установлен пакет xlrd. Выполните: pip install -r requirements.txt'
        ) from exc
    book = xlrd.open_workbook(str(path), on_demand=True)
    try:
        sh = book.sheet_by_index(0)
        return [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
    finally:
        book.release_resources()


def read_table_rows(path: str | Path) -> list[list]:
    suffix = Path(path).suffix.casefold()
    if suffix == '.xlsx':
        return _read_xlsx_rows(path)
    if suffix == '.xls':
        return _read_xls_rows(path)
    raise ValueError('Поддерживаются только файлы .xls и .xlsx')


def load_stock_file(path: str | Path) -> tuple[dict[str, int], list[tuple[str, str, str]]]:
    rows = read_table_rows(path)
    errors: list[tuple[str, str, str]] = []
    header_idx = code_col = qty_col = None

    for r_idx, row in enumerate(rows[:30]):
        headers = [_norm_header(x) for x in row]
        for c_idx, h in enumerate(headers):
            if h in {'код', 'баркод', 'barcode', 'sku'}:
                code_col = c_idx
            if h in {'доступно', 'количество', 'остаток', 'qty', 'quantity'}:
                qty_col = c_idx
        if code_col is not None and qty_col is not None:
            header_idx = r_idx
            break
        code_col = qty_col = None

    if header_idx is None or code_col is None or qty_col is None:
        raise ValueError('В файле остатков не найдены колонки «Код» и «Доступно».')

    stocks: dict[str, int] = defaultdict(int)
    for row_num, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if code_col >= len(row):
            continue
        bc = barcode_text(row[code_col])
        if not bc:
            continue
        qty = int_qty(row[qty_col] if qty_col < len(row) else 0)
        stocks[bc] += qty
        if qty < 0:
            errors.append(('Остатки', bc, f'Строка {row_num}: отрицательный остаток {qty}.'))
    return dict(stocks), errors


def load_bundles_file(path: str | Path) -> tuple[list[Bundle], list[tuple[str, str, str]]]:
    rows = read_table_rows(path)
    errors: list[tuple[str, str, str]] = []
    if not rows:
        raise ValueError('Шаблон комплектов пуст.')

    header_idx = name_col = bundle_col = None
    component_cols: list[int] = []
    for r_idx, row in enumerate(rows[:30]):
        headers = [_norm_header(x) for x in row]
        name_col = next((i for i, h in enumerate(headers) if h in {'название', 'наименование'}), None)
        bundle_col = next((i for i, h in enumerate(headers) if 'баркод' in h and ('комплект' in h or 'набор' in h)), None)
        if bundle_col is None:
            bundle_col = next((i for i, h in enumerate(headers) if h in {'баркодкомплекта', 'баркоднабора'}), None)
        if bundle_col is not None:
            component_cols = [
                i for i, h in enumerate(headers)
                if i != bundle_col and (h.startswith('баркод') or h.startswith('barcode'))
            ]
        if bundle_col is not None and component_cols:
            header_idx = r_idx
            break

    if header_idx is None or bundle_col is None or not component_cols:
        raise ValueError(
            'В шаблоне комплектов нужны колонки «баркод комплекта» и «баркод1», «баркод2» ...'
        )

    bundles: list[Bundle] = []
    seen_bundle_barcodes: set[str] = set()
    for row_num, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        bundle_bc = barcode_text(row[bundle_col] if bundle_col < len(row) else None)
        if not bundle_bc:
            continue
        name = ''
        if name_col is not None and name_col < len(row):
            name = str(row[name_col] or '').strip()
        components = [
            barcode_text(row[c] if c < len(row) else None)
            for c in component_cols
        ]
        components = [x for x in components if x]
        if not components:
            errors.append(('Комплекты', bundle_bc, f'Строка {row_num}: у комплекта нет компонентов.'))
            continue
        if bundle_bc in seen_bundle_barcodes:
            errors.append(('Комплекты', bundle_bc, f'Строка {row_num}: повторный баркод комплекта.'))
            continue
        seen_bundle_barcodes.add(bundle_bc)
        bundles.append(Bundle(name=name or bundle_bc, barcode=bundle_bc, components=components))
    return bundles, errors


def _product_barcode_index(products: Iterable[Product]) -> dict[str, Product]:
    index: dict[str, Product] = {}
    for p in products:
        for bc in p.barcodes:
            if bc:
                index.setdefault(bc, p)
    return index


def calculate_procurement(
    products: list[Product],
    stocks_fbs: dict[str, int],
    bundles: list[Bundle],
    analysis_days: int,
    target_days: int,
    coefficient: float,
    stock_mode: str = 'fbs',
    initial_errors: list[tuple[str, str, str]] | None = None,
) -> ProcurementResult:
    errors = list(initial_errors or [])
    warnings = []
    by_barcode = _product_barcode_index(products)
    bundle_barcodes = {b.barcode for b in bundles}

    # Products that are bundle SKUs are shown on the audit sheet, not in the purchase list.
    bundle_products = {id(by_barcode[bc]) for bc in bundle_barcodes if bc in by_barcode}

    fbs_direct_by_product: dict[int, int] = defaultdict(int)
    fbo_direct_by_product: dict[int, int] = defaultdict(int)
    for p in products:
        file_qty = sum(stocks_fbs.get(bc, 0) for bc in p.barcodes)
        fbs_direct_by_product[id(p)] = file_qty if stocks_fbs else int(p.fbs_stock or 0)
        fbo_direct_by_product[id(p)] = int(p.fbo_stock or 0)

    bundle_sales_by_product: dict[int, int] = defaultdict(int)
    bundle_rate_by_product: dict[int, Decimal] = defaultdict(Decimal)
    fbs_stock_in_bundles_by_product: dict[int, int] = defaultdict(int)
    fbo_stock_in_bundles_by_product: dict[int, int] = defaultdict(int)
    bundle_rows: list[BundleAuditRow] = []

    def rate(product):
        days = len(product.availability_dates) if product.availability_dates is not None else analysis_days
        if days == 0 and product.total > 0:
            days = analysis_days
        return Decimal(product.total) / Decimal(days) if days > 0 else Decimal(0)

    def selected_stock(fbo: int, fbs: int) -> int:
        if stock_mode == 'fbo':
            return fbo
        if stock_mode == 'both':
            return fbo + fbs
        return fbs

    for bundle in bundles:
        bundle_product = by_barcode.get(bundle.barcode)
        if bundle_product is None:
            errors.append(('Комплекты', bundle.barcode, 'Баркод комплекта не найден среди товаров Wildberries.'))
            fbo = fbs = 0
            bundle_stock_fbo = 0
        else:
            fbo, fbs = bundle_product.fbo, bundle_product.fbs
            bundle_stock_fbo = int(bundle_product.fbo_stock or 0)
        bundle_orders = fbo + fbs
        bundle_stock_fbs = (
            stocks_fbs.get(bundle.barcode, 0)
            if stocks_fbs
            else int(bundle_product.fbs_stock or 0) if bundle_product is not None else 0
        )
        bundle_stock_selected = selected_stock(bundle_stock_fbo, bundle_stock_fbs)

        matched_counts: dict[int, int] = defaultdict(int)
        matched_product: dict[int, Product] = {}
        representative_bc: dict[int, str] = {}
        unknown_counts: Counter[str] = Counter()
        for comp_bc in bundle.components:
            comp_product = by_barcode.get(comp_bc)
            if comp_product is None:
                unknown_counts[comp_bc] += 1
                continue
            pid = id(comp_product)
            matched_counts[pid] += 1
            matched_product[pid] = comp_product
            representative_bc.setdefault(pid, comp_bc)

        for comp_bc, mult in unknown_counts.items():
            errors.append((
                'Комплекты', comp_bc,
                f'Компонент не найден среди товаров WB. Комплект: {bundle.name} ({bundle.barcode}), кратность {mult}.'
            ))
            bundle_rows.append(BundleAuditRow(
                bundle.name, bundle.barcode, fbo, fbs, bundle_orders,
                bundle_stock_fbo, bundle_stock_fbs, bundle_stock_selected,
                comp_bc, 'НЕ НАЙДЕН', mult, bundle_orders * mult,
                bundle_stock_fbo * mult, bundle_stock_fbs * mult, bundle_stock_selected * mult,
            ))

        for pid, mult in matched_counts.items():
            p = matched_product[pid]
            if pid in bundle_products:
                errors.append((
                    'Комплекты', representative_bc[pid],
                    f'Компонент сам является комплектом ({p.name}). Вложенные комплекты не разворачиваются автоматически.'
                ))
            sales_use = bundle_orders * mult
            stock_use_fbo = bundle_stock_fbo * mult
            stock_use_fbs = bundle_stock_fbs * mult
            stock_use_selected = bundle_stock_selected * mult
            bundle_sales_by_product[pid] += sales_use
            if bundle_product is not None:
                bundle_rate_by_product[pid] += rate(bundle_product) * mult
            fbo_stock_in_bundles_by_product[pid] += stock_use_fbo
            fbs_stock_in_bundles_by_product[pid] += stock_use_fbs
            bundle_rows.append(BundleAuditRow(
                bundle.name, bundle.barcode, fbo, fbs, bundle_orders,
                bundle_stock_fbo, bundle_stock_fbs, bundle_stock_selected,
                representative_bc[pid], p.name, mult, sales_use,
                stock_use_fbo, stock_use_fbs, stock_use_selected,
            ))

    rows: list[PurchaseRow] = []
    for p in products:
        if id(p) in bundle_products:
            continue
        row = PurchaseRow(product=p)
        row.bundle_sales_units = bundle_sales_by_product[id(p)]
        row.fbo_direct_stock = fbo_direct_by_product[id(p)]
        row.fbs_direct_stock = fbs_direct_by_product[id(p)]
        row.fbo_stock_inside_bundles = fbo_stock_in_bundles_by_product[id(p)]
        row.fbs_stock_inside_bundles = fbs_stock_in_bundles_by_product[id(p)]
        row.selected_stock = selected_stock(row.fbo_physical_stock, row.fbs_physical_stock)
        total = row.total_consumption
        daily_rate = rate(p) + bundle_rate_by_product[id(p)]
        row.average_per_day = float(daily_rate)
        raw_required = daily_rate * Decimal(target_days) * Decimal(str(coefficient))
        row.required_stock = int(raw_required.to_integral_value(rounding=ROUND_CEILING))
        row.purchase_qty = max(0, row.required_stock - row.physical_stock)
        if stocks_fbs and stock_mode in {'fbs', 'both'} and not any(bc in stocks_fbs for bc in p.barcodes) and row.fbs_stock_inside_bundles == 0:
            row.note = 'Нет в файле остатков FBS'
        if p.history_warning:
            row.note = '; '.join(filter(None, (row.note, p.history_warning)))
        if p.availability_dates is not None and not p.availability_dates and p.total > 0:
            row.note = '; '.join(filter(None, (row.note, 'Нет подтверждённых дней наличия при заказах: расчёт по календарным дням.')))
        row.cabinet_days[p.cabinet] = len(p.availability_dates) if p.availability_dates is not None else None
        for bundle in bundles:
            if any(by_barcode.get(barcode) is p for barcode in bundle.components):
                bundle_product = by_barcode.get(bundle.barcode)
                if bundle_product is not None and bundle_product.history_warning:
                    row.note = '; '.join(filter(None, (row.note, f'Комплект {bundle.barcode}: {bundle_product.history_warning}')))
        rows.append(row)

    known_barcodes = set(by_barcode) | bundle_barcodes
    for bc, qty in stocks_fbs.items():
        if bc not in known_barcodes:
            errors.append(('Остатки', bc, f'Баркод из файла остатков не найден в WB/шаблоне комплектов. Остаток: {qty}.'))

    if any(product.cabinet_products for product in products):
        separate: dict[str, list[Product]] = defaultdict(list)
        for product in products:
            for member in product.cabinet_products or [product]:
                separate[member.cabinet].append(replace(member, barcodes=set(product.barcodes), cabinet_products=[]))
        for cabinet, members in separate.items():
            cabinet_result = calculate_procurement(members, {}, bundles, analysis_days, target_days, coefficient, 'fbo')
            for row in rows:
                member_ids = {(member.nm_id, member.chrt_id) for member in row.product.cabinet_products or [row.product] if member.cabinet == cabinet}
                matches = [candidate for candidate in cabinet_result.rows if (candidate.product.nm_id, candidate.product.chrt_id) in member_ids]
                row.cabinet_fbo[cabinet] = sum(candidate.fbo_physical_stock for candidate in matches)
                row.cabinet_required[cabinet] = sum(candidate.required_stock for candidate in matches)
                row.cabinet_needs[cabinet] = sum(
                    max(0, candidate.required_stock - (candidate.fbo_physical_stock if stock_mode in {'fbo', 'both'} else 0))
                    for candidate in matches
                )
                row.cabinet_average[cabinet] = sum(candidate.average_per_day for candidate in matches)
                days = [candidate.cabinet_days.get(cabinet) for candidate in matches]
                row.cabinet_days[cabinet] = days[0] if len(days) == 1 else None
                notes = [row.note] + [f'{cabinet}: {candidate.note}' for candidate in matches if candidate.note]
                row.note = '; '.join(dict.fromkeys(note for note in notes if note))
        for row in rows:
            row.average_per_day = sum(row.cabinet_average.values())
            row.required_stock = sum(row.cabinet_required.values())
            row.total_need = sum(row.cabinet_needs.values())
            row.purchase_qty = max(0, row.total_need - (row.fbs_physical_stock if stock_mode in {'fbs', 'both'} else 0))
    else:
        for row in rows:
            cabinet = row.product.cabinet
            row.cabinet_fbo[cabinet] = row.fbo_physical_stock
            row.cabinet_average[cabinet] = row.average_per_day
            row.cabinet_required[cabinet] = row.required_stock
            row.total_need = max(0, row.required_stock - (row.fbo_physical_stock if stock_mode in {'fbo', 'both'} else 0))
            row.cabinet_needs[cabinet] = row.total_need
    rows.sort(key=lambda r: (-r.purchase_qty, r.product.name.casefold(), r.product.size.casefold()))
    bundle_rows.sort(key=lambda r: (r.bundle_name.casefold(), r.component_name.casefold()))
    for product in products:
        for member in product.cabinet_products or [product]:
            if member.history_warning:
                target = warnings if member.history_status in {'not_returned', 'missing_values', 'incomplete_period'} else errors
                target.append(('История наличия', ', '.join(sorted(member.barcodes)), f'{member.cabinet}: {member.history_warning}'))
    return ProcurementResult(rows=rows, bundle_rows=bundle_rows, errors=errors, warnings=warnings, source_products=[member for product in products for member in product.cabinet_products or [product]])


def make_procurement_excel(
    result: ProcurementResult,
    path: str | Path,
    cabinet_name: str,
    period_text: str,
    analysis_days: int,
    target_days: int,
    coefficient: float,
    stock_mode: str = 'fbs',
):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Закупка'

    navy = '1F4E78'
    blue = 'D9EAF7'
    green = 'E2F0D9'
    yellow = 'FFF2CC'
    red = 'FCE4D6'
    thin = Side(style='thin', color='D9E2F3')

    cabinets = list(dict.fromkeys(cabinet for row in result.rows for cabinet in row.cabinet_needs))
    last_column = get_column_letter(20 + 5 * len(cabinets))
    ws.merge_cells(f'A1:{last_column}1')
    ws['A1'] = 'Расчёт закупки Wildberries'
    ws['A1'].font = Font(bold=True, size=16, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor=navy)
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells(f'A2:{last_column}2')
    ws['A2'] = f'Кабинет: {cabinet_name} | Период продаж: {period_text} ({analysis_days} дн.)'
    ws.merge_cells(f'A3:{last_column}3')
    mode_label = {'fbo': 'Только FBO', 'fbs': 'Только FBS', 'both': 'FBO + FBS'}.get(stock_mode, stock_mode)
    ws['A3'] = f'Целевой запас: {target_days} дн. | Коэффициент: {coefficient:g} | Учитываемые остатки: {mode_label}'

    headers = [
        'Баркод(ы)', 'Артикул продавца', 'Бренд', 'Наименование', 'К закупке', 'На сколько хватит дней',
        'Размер', 'Заказы FBO', 'Заказы FBS', 'Продажи комплектами', 'Общий расход', 'Среднее/день',
        'Запас, дней', 'Коэффициент', 'Необходимо иметь', 'Остаток FBO', 'Остаток FBS', 'Остаток учтён', 'Примечание'
    ]
    headers += [label for cabinet in cabinets for label in (f'Остаток FBO — {cabinet}', f'Целевой запас — {cabinet}', f'Потребность — {cabinet}', f'Дней наличия — {cabinet}', f'Среднее/день — {cabinet}')]
    headers += ['Потребность общая до FBS']
    header_row = 5
    for col, value in enumerate(headers, 1):
        c = ws.cell(header_row, col, value)
        c.fill = PatternFill('solid', fgColor=navy)
        c.font = Font(color='FFFFFF', bold=True)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for r_idx, r in enumerate(result.rows, start=header_row + 1):
        values = [
            ', '.join(sorted(r.product.barcodes)), r.product.vendor_code, r.product.brand, r.product.name,
            r.purchase_qty, r.days_remaining, r.product.size or '0', r.product.fbo, r.product.fbs,
            r.bundle_sales_units, r.total_consumption, r.average_per_day, target_days, coefficient,
            r.required_stock, r.fbo_physical_stock, r.fbs_physical_stock, r.physical_stock, r.note,
        ]
        values += [value for cabinet in cabinets for value in (r.cabinet_fbo.get(cabinet, 0), r.cabinet_required.get(cabinet, 0), r.cabinet_needs.get(cabinet, 0), r.cabinet_days.get(cabinet) if r.cabinet_days.get(cabinet) is not None else 'Нет полной истории', r.cabinet_average.get(cabinet, 0))]
        values += [r.total_need]
        for c_idx, value in enumerate(values, 1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical='center', wrap_text=c_idx in (1, 4, 19))
        ws.cell(r_idx, 6).number_format = '0.00'
        ws.cell(r_idx, 12).number_format = '0.00'
        ws.cell(r_idx, 14).number_format = '0.00'
        if r.purchase_qty > 0:
            ws.cell(r_idx, 5).fill = PatternFill('solid', fgColor=green)
            ws.cell(r_idx, 5).font = Font(bold=True)
        if r.note:
            ws.cell(r_idx, 19).fill = PatternFill('solid', fgColor=yellow)

    widths = [32, 24, 20, 48, 14, 20, 12, 13, 13, 20, 14, 14, 12, 13, 18, 14, 14, 16, 26]
    widths += [24] * (5 * len(cabinets) + 1)
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'A6'
    if result.rows:
        ws.auto_filter.ref = f'A5:{last_column}{header_row + len(result.rows)}'
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[5].height = 36

    # Bundle audit sheet
    bs = wb.create_sheet('Комплекты')
    b_headers = [
        'Название', 'Баркод комплекта', 'Заказы FBO', 'Заказы FBS', 'Всего заказов',
        'Остаток комплектов FBO', 'Остаток комплектов FBS', 'Остаток комплектов учтён',
        'Баркод компонента', 'Наименование компонента', 'Кол-во в комплекте', 'Расход через продажи',
        'Компонентов в FBO-комплектах', 'Компонентов в FBS-комплектах', 'Компонентов учтено'
    ]
    for i, h in enumerate(b_headers, 1):
        c = bs.cell(1, i, h)
        c.fill = PatternFill('solid', fgColor=navy)
        c.font = Font(color='FFFFFF', bold=True)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r_idx, r in enumerate(result.bundle_rows, 2):
        vals = [
            r.bundle_name, r.bundle_barcode, r.fbo, r.fbs, r.bundle_orders,
            r.bundle_stock_fbo, r.bundle_stock_fbs, r.bundle_stock_selected,
            r.component_barcode, r.component_name, r.qty_in_bundle, r.component_sales_consumption,
            r.component_stock_inside_bundles_fbo, r.component_stock_inside_bundles_fbs, r.component_stock_inside_bundles_selected,
        ]
        for c_idx, value in enumerate(vals, 1):
            cell = bs.cell(r_idx, c_idx, value)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical='center', wrap_text=c_idx in (1, 10))
    for i, width in enumerate([24, 22, 12, 12, 14, 18, 18, 20, 22, 42, 18, 20, 24, 24, 22], 1):
        bs.column_dimensions[get_column_letter(i)].width = width
    bs.freeze_panes = 'A2'
    if result.bundle_rows:
        bs.auto_filter.ref = f'A1:O{1 + len(result.bundle_rows)}'
    bs.row_dimensions[1].height = 42

    es = wb.create_sheet('Ошибки')
    for i, h in enumerate(['Источник', 'Баркод', 'Описание'], 1):
        c = es.cell(1, i, h)
        c.fill = PatternFill('solid', fgColor=navy)
        c.font = Font(color='FFFFFF', bold=True)
    if result.errors:
        for r_idx, (source, bc, desc) in enumerate(result.errors, 2):
            es.cell(r_idx, 1, source)
            es.cell(r_idx, 2, bc)
            es.cell(r_idx, 3, desc)
            for c_idx in range(1, 4):
                es.cell(r_idx, c_idx).fill = PatternFill('solid', fgColor=red if source == 'Комплекты' else yellow)
                es.cell(r_idx, c_idx).border = Border(left=thin, right=thin, top=thin, bottom=thin)
                es.cell(r_idx, c_idx).alignment = Alignment(vertical='top', wrap_text=True)
    else:
        es['A2'] = 'Ошибок не обнаружено'
        es['A2'].fill = PatternFill('solid', fgColor=green)
    es.column_dimensions['A'].width = 18
    es.column_dimensions['B'].width = 24
    es.column_dimensions['C'].width = 90
    es.freeze_panes = 'A2'

    warning_sheet = wb.create_sheet('Предупреждения')
    warning_sheet.append(['Источник', 'Баркод', 'Описание'])
    for cell in warning_sheet[1]:
        cell.fill = PatternFill('solid', fgColor=navy)
        cell.font = Font(color='FFFFFF', bold=True)
    for warning in result.warnings:
        warning_sheet.append(warning)
        for cell in warning_sheet[warning_sheet.max_row]:
            cell.fill = PatternFill('solid', fgColor=yellow)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    if not result.warnings:
        warning_sheet.append(['Предупреждений нет'])
    warning_sheet.column_dimensions['A'].width = 22
    warning_sheet.column_dimensions['B'].width = 32
    warning_sheet.column_dimensions['C'].width = 100
    warning_sheet.freeze_panes = 'A2'

    availability_sheet = wb.create_sheet('Наличие')
    availability_sheet.append(['Кабинет', 'Артикул WB', 'ID размера', 'Баркод(ы)', 'Наименование', 'Размер', 'Заказы напрямую', 'Дней анализа', 'Дней наличия напрямую', 'Среднее прямых заказов/день', 'Основание расчёта'])
    for cell in availability_sheet[1]:
        cell.fill = PatternFill('solid', fgColor=navy)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for product in result.source_products:
        active_days = len(product.availability_dates) if product.availability_dates is not None else None
        denominator = active_days if active_days is not None else analysis_days
        if denominator == 0 and product.total > 0:
            denominator = analysis_days
        average = product.total / denominator if denominator > 0 else 0
        availability_label = 'FBS' if product.availability_mode == 'fbs' else 'FBO/FBS'
        basis = product.history_warning or (f'История {availability_label} на конец дня + дни заказов {availability_label}' if active_days is not None else f'Календарные дни: нет полной истории {availability_label}')
        availability_sheet.append([product.cabinet, product.nm_id, product.chrt_id, ', '.join(sorted(product.barcodes)), product.name, product.size, product.total, analysis_days, active_days, average, basis])
        availability_sheet.cell(availability_sheet.max_row, 10).number_format = '0.00'
    for index, width in enumerate([24, 18, 18, 36, 48, 14, 18, 16, 24, 24, 90], 1):
        availability_sheet.column_dimensions[get_column_letter(index)].width = width
    availability_sheet.freeze_panes = 'A2'
    availability_sheet.auto_filter.ref = availability_sheet.dimensions
    availability_sheet.row_dimensions[1].height = 36

    wb.save(path)
