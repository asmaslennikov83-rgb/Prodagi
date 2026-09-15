from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


@dataclass
class Product:
    cabinet: str
    nm_id: int
    chrt_id: int
    vendor_code: str
    name: str
    size: str
    barcodes: set[str] = field(default_factory=set)
    fbo: int = 0
    fbs: int = 0

    @property
    def total(self) -> int:
        return self.fbo + self.fbs


def s(v) -> str:
    return '' if v is None else str(v).strip()


def build_products(cards: list[dict], cabinet: str) -> list[Product]:
    products: list[Product] = []
    for card in cards:
        nm_id = int(card.get('nmID') or 0)
        # Seller article is normally vendorCode in Content API. Keep fallbacks for compatibility.
        vendor = s(card.get('vendorCode') or card.get('supplierArticle') or card.get('vendor_code'))
        name = s(card.get('title'))
        sizes = card.get('sizes') or []
        if not sizes:
            products.append(Product(cabinet, nm_id, 0, vendor, name, '0', set()))
            continue
        for size in sizes:
            chrt_id = int(size.get('chrtID') or 0)
            tech_size = s(size.get('techSize') or size.get('wbSize')) or '0'
            barcodes = {s(x) for x in (size.get('skus') or []) if s(x)}
            products.append(Product(cabinet, nm_id, chrt_id, vendor, name, tech_size, barcodes))
    return products


def is_cancelled(order: dict, source: str) -> bool:
    if source == 'orders-api':
        return bool(order.get('isCancel'))
    return s(order.get('status')).lower() in {'cancel', 'canceled', 'cancelled'}


def order_model(order: dict, source: str) -> str | None:
    if source == 'orders-api':
        wt = s(order.get('warehouseType')).lower()
        if 'продав' in wt or 'seller' in wt or 'fbs' in wt:
            return 'FBS'
        return 'FBO'
    # Order Feed: MP = marketplace/seller warehouse, otherwise WB warehouse.
    return 'FBS' if bool(order.get('isMp')) else 'FBO'


def apply_orders(products: list[Product], orders: list[dict], source: str) -> int:
    by_barcode: dict[str, Product] = {}
    by_chrt: dict[int, Product] = {}
    by_nm: dict[int, list[Product]] = defaultdict(list)
    for p in products:
        for bc in p.barcodes:
            by_barcode[bc] = p
        if p.chrt_id:
            by_chrt[p.chrt_id] = p
        if p.nm_id:
            by_nm[p.nm_id].append(p)

    unmatched = 0
    seen: set[str] = set()
    for o in orders:
        if is_cancelled(o, source):
            continue
        srid = s(o.get('srid'))
        if srid and srid in seen:
            continue
        if srid:
            seen.add(srid)

        p = None
        bc = s(o.get('barcode'))
        if bc:
            p = by_barcode.get(bc)
        if p is None:
            try:
                chrt = int(o.get('chrtId') or o.get('chrtID') or 0)
            except Exception:
                chrt = 0
            if chrt:
                p = by_chrt.get(chrt)
        if p is None:
            try:
                nm = int(o.get('nmId') or o.get('nmID') or 0)
            except Exception:
                nm = 0
            candidates = by_nm.get(nm, [])
            if len(candidates) == 1:
                p = candidates[0]
            elif candidates:
                size = s(o.get('techSize') or o.get('size')) or '0'
                p = next((x for x in candidates if x.size == size), None)

        if p is None:
            unmatched += 1
            continue

        # Orders API also contains seller article (supplierArticle).
        # Use it as a fallback if Content API returned an empty vendorCode.
        if not p.vendor_code:
            p.vendor_code = s(o.get('supplierArticle') or o.get('vendorCode') or o.get('vendor_code'))

        model = order_model(o, source)
        if model == 'FBS':
            p.fbs += 1
        elif model == 'FBO':
            p.fbo += 1
    return unmatched


class DSU:
    def __init__(self, n: int):
        self.p = list(range(n))
    def find(self, x: int) -> int:
        if self.p[x] != x:
            self.p[x] = self.find(self.p[x])
        return self.p[x]
    def union(self, a: int, b: int):
        a, b = self.find(a), self.find(b)
        if a != b:
            self.p[b] = a


def merge_products(cab1: list[Product], cab2: list[Product]) -> list[Product]:
    """Merge primarily by ANY intersecting barcode across cabinets.
    Within the same cabinet, keep sizes separate. If cross-cabinet barcodes don't intersect,
    fallback to vendorCode + normalized size to avoid losing identical listings with changed barcodes.
    """
    items = cab1 + cab2
    dsu = DSU(len(items))
    barcode_owner: dict[str, int] = {}
    for i, p in enumerate(items):
        for bc in p.barcodes:
            if bc in barcode_owner:
                j = barcode_owner[bc]
                if items[j].cabinet != p.cabinet:
                    dsu.union(i, j)
            else:
                barcode_owner[bc] = i

    logical_owner: dict[tuple[str, str], int] = {}
    for i, p in enumerate(items):
        key = (p.vendor_code.casefold(), p.size.casefold())
        if not key[0]:
            continue
        if key in logical_owner:
            j = logical_owner[key]
            if items[j].cabinet != p.cabinet:
                dsu.union(i, j)
        else:
            logical_owner[key] = i

    groups: dict[int, list[Product]] = defaultdict(list)
    for i, p in enumerate(items):
        groups[dsu.find(i)].append(p)

    out: list[Product] = []
    for group in groups.values():
        first = group[0]
        # Do not lose seller article/name when the first matched card has an empty field.
        vendor_code = next((g.vendor_code for g in group if g.vendor_code), '')
        name = next((g.name for g in group if g.name), '')
        nm_id = next((g.nm_id for g in group if g.nm_id), 0)
        chrt_id = next((g.chrt_id for g in group if g.chrt_id), 0)
        size = next((g.size for g in group if g.size and g.size != '0'), first.size or '0')
        out.append(Product(
            cabinet='Оба кабинета' if len({g.cabinet for g in group}) > 1 else first.cabinet,
            nm_id=nm_id,
            chrt_id=chrt_id,
            vendor_code=vendor_code,
            name=name,
            size=size or '0',
            barcodes=set().union(*(g.barcodes for g in group)),
            fbo=sum(g.fbo for g in group),
            fbs=sum(g.fbs for g in group),
        ))
    return out


def make_excel(products: Iterable[Product], path: str | Path, title: str, period_text: str):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Заказы'
    ws.append(['Баркод(ы)', 'Артикул продавца', 'nmID', 'Наименование', 'Размер', 'Заказы FBO', 'Заказы FBS', 'Всего'])

    header_fill = PatternFill('solid', fgColor='1F4E78')
    header_font = Font(color='FFFFFF', bold=True)
    thin = Side(style='thin', color='D9E2F3')
    for c in ws[1]:
        c.fill, c.font = header_fill, header_font
        c.alignment = Alignment(horizontal='center', vertical='center')

    for p in sorted(products, key=lambda x: (x.name.casefold(), x.vendor_code.casefold(), x.size.casefold())):
        ws.append([', '.join(sorted(p.barcodes)), p.vendor_code, p.nm_id, p.name, p.size or '0', p.fbo, p.fbs, p.total])

    for row in ws.iter_rows():
        for c in row:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            c.alignment = Alignment(vertical='center')

    for col, width in {'A': 32, 'B': 25, 'C': 15, 'D': 52, 'E': 14, 'F': 15, 'G': 15, 'H': 15}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet('Информация')
    info.append(['Отчёт', title])
    info.append(['Период', period_text])
    info.append(['Примечание', 'Сегодняшний день не включается в быстрые периоды. Отменённые заказы исключены.'])
    info.column_dimensions['A'].width = 22
    info.column_dimensions['B'].width = 90
    wb.save(path)
